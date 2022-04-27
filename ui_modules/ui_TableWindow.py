from PyQt5 import QtGui
from PyQt5.QtWidgets import QDialog, QHeaderView, QAbstractItemView

from ui_modules.sampleWindow import SampleDialog
from ui_modules.ui.table import UiTableWindow

import docx


class TableWindow(SampleDialog, UiTableWindow):
    def __init__(self, parent=None):
        QDialog.__init__(self)
        self.setupUi(self)
        self.titleWindowLoad()

        self.data = None

        self.model = QtGui.QStandardItemModel()
        self.tableView.setModel(self.model)
        self.tableView.verticalHeader().setVisible(False)
        self.tableView.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeToContents)
        self.tableView.verticalHeader().setSectionResizeMode(QHeaderView.ResizeToContents)
        self.tableView.setEditTriggers(QAbstractItemView.NoEditTriggers)

        self.pushButton_save.clicked.connect(self.__toWord)

    def setData(self, data: list):
        self.data = data
        rowCount = len(self.data)

        if rowCount == 0:
            return 0

        self.model.clear()
        self.model.setHorizontalHeaderLabels(
            ['Код\nоперации', 'Название', 'Откат к', 'Примечание', 'Выполняющий', 'Дата', 'Время'])

        for i in self.data:
            i = [QtGui.QStandardItem(i) for i in i]
            self.model.appendRow(i)

        self.showWind()

    def __toWord(self):
        row = len(self.data)
        if self.data is None or row == 0:
            return 0

        data = [['Код операции', 'Название', 'Откат к', 'Примечание', 'Выполняющий', 'Дата', 'Время']] + self.data
        doc = docx.Document()

        table = doc.add_table(rows=row, cols=7)
        table.style = 'Table Grid'

        for row in range(row):
            for col in range(7):
                cell = table.cell(row, col)
                cell.text = data[row][col]

        doc.save('table.docx')
